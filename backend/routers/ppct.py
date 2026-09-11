from fastapi import APIRouter, UploadFile, File, Form, Depends, HTTPException, Query
from starlette.concurrency import run_in_threadpool
from sqlalchemy.orm import Session
from typing import List, Optional

from db.session import get_db
from db.models import PPCTModel
from schemas.ppct import PPCTItemResponse, PPCTItemCreate, PPCTItemUpdate, PPCTBulkUploadResponse
from services.ppct_parser import parse_ppct_excel, parse_ppct_word, parse_ppct_pdf
from routers.auth import get_current_user_id
from logger import logger

router = APIRouter(prefix="/ppct", tags=["Phân Phối Chương Trình (PPCT)"])

def normalize_grade(grade_input: Optional[str], context_hint: str = "") -> str:
    g = str(grade_input or "").strip().lower()
    h = context_hint.lower()
    if "12" in g or "khối 12" in h or "lớp 12" in h or "k12" in h or "toán 12" in h or "12" in h:
        return "Khối 12"
    if "11" in g or "khối 11" in h or "lớp 11" in h or "k11" in h or "toán 11" in h or "11" in h:
        return "Khối 11"
    if "10" in g or "khối 10" in h or "lớp 10" in h or "k10" in h or "toán 10" in h or "10" in h:
        return "Khối 10"
    return grade_input.strip() if (grade_input and grade_input.strip()) else "Khối 10"

@router.post("/upload", response_model=PPCTBulkUploadResponse)
async def upload_ppct_file(
    file: UploadFile = File(...),
    grade: str = Form("Khối 10"),
    subject: str = Form("Toán"),
    overwrite: bool = Form(True),
    api_key: Optional[str] = Form(None),
    model_name: Optional[str] = Form(None),
    user_id: str = Depends(get_current_user_id),
    db: Session = Depends(get_db)
):
    from services.gemini_service import set_custom_gemini_key
    if api_key and api_key.strip():
        await run_in_threadpool(set_custom_gemini_key, api_key.strip())

    content = await file.read()
    filename = file.filename.lower()
    
    # Auto-detect / normalize grade (e.g. Khối 12)
    grade = normalize_grade(grade, filename)
    
    parsed_items = []
    if filename.endswith(".xlsx") or filename.endswith(".xls"):
        parsed_items = await run_in_threadpool(parse_ppct_excel, content, default_grade=grade, default_subject=subject)
    elif filename.endswith(".docx"):
        parsed_items = await run_in_threadpool(parse_ppct_word, content, default_grade=grade, default_subject=subject)
    elif filename.endswith(".pdf"):
        parsed_items = await run_in_threadpool(parse_ppct_pdf, content, default_grade=grade, default_subject=subject, selected_model=model_name)
    else:
        raise HTTPException(status_code=400, detail="Chỉ hỗ trợ file định dạng Excel (.xlsx), Word (.docx) hoặc PDF (.pdf)")

    if not parsed_items:
        raise HTTPException(status_code=400, detail="Không thể đọc nội dung phân phối chương trình từ file. Vui lòng kiểm tra định dạng bảng.")

    # Overwrite previous items for this user & grade & subject if requested
    if overwrite:
        db.query(PPCTModel).filter(
            PPCTModel.user_id == user_id,
            PPCTModel.grade == grade,
            PPCTModel.subject == subject
        ).delete()
        db.commit()

    saved_items = []
    for item_data in parsed_items:
        ppct_item = PPCTModel(
            user_id=user_id,
            grade=item_data.get("grade", grade),
            subject=item_data.get("subject", subject),
            week=item_data.get("week", 1),
            lesson_number=item_data.get("lesson_number", 1),
            lesson_title=item_data.get("lesson_title", "Bài học"),
            notes=item_data.get("notes"),
            semester=item_data.get("semester", "Học kỳ 1")
        )
        db.add(ppct_item)
        saved_items.append(ppct_item)

    db.commit()
    for item in saved_items:
        db.refresh(item)

    return PPCTBulkUploadResponse(
        total_imported=len(saved_items),
        grade=grade,
        subject=subject,
        items=saved_items
    )

@router.get("/", response_model=List[PPCTItemResponse])
async def list_ppct(
    grade: Optional[str] = Query(None),
    subject: Optional[str] = Query(None),
    week: Optional[int] = Query(None),
    user_id: str = Depends(get_current_user_id),
    db: Session = Depends(get_db)
):
    q = db.query(PPCTModel).filter(PPCTModel.user_id == user_id)
    if grade:
        q = q.filter(PPCTModel.grade == grade)
    if subject:
        q = q.filter(PPCTModel.subject == subject)
    if week:
        q = q.filter(PPCTModel.week == week)
    return q.order_by(PPCTModel.week.asc(), PPCTModel.lesson_number.asc()).all()

@router.post("/", response_model=PPCTItemResponse)
async def create_ppct_item(
    item_in: PPCTItemCreate,
    user_id: str = Depends(get_current_user_id),
    db: Session = Depends(get_db)
):
    db_item = PPCTModel(**item_in.model_dump(), user_id=user_id)
    db.add(db_item)
    db.commit()
    db.refresh(db_item)
    return db_item

@router.put("/{item_id}", response_model=PPCTItemResponse)
async def update_ppct_item(
    item_id: int,
    item_in: PPCTItemUpdate,
    user_id: str = Depends(get_current_user_id),
    db: Session = Depends(get_db)
):
    item = db.query(PPCTModel).filter(PPCTModel.id == item_id, PPCTModel.user_id == user_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Bản ghi không tồn tại")
    
    update_data = item_in.model_dump(exclude_unset=True)
    for k, v in update_data.items():
        if v is not None:
            setattr(item, k, v)
            
    db.commit()
    db.refresh(item)
    return item

@router.get("/stats")
async def get_ppct_stats(
    user_id: str = Depends(get_current_user_id),
    db: Session = Depends(get_db)
):
    total = db.query(PPCTModel).filter(PPCTModel.user_id == user_id).count()
    from sqlalchemy import func
    grade_counts = db.query(PPCTModel.grade, func.count(PPCTModel.id)).filter(PPCTModel.user_id == user_id).group_by(PPCTModel.grade).all()
    subject_counts = db.query(PPCTModel.subject, func.count(PPCTModel.id)).filter(PPCTModel.user_id == user_id).group_by(PPCTModel.subject).all()
    return {
        "total": total,
        "grades": [{"grade": g, "count": c} for g, c in grade_counts],
        "subjects": [{"subject": s, "count": c} for s, c in subject_counts]
    }

@router.delete("/clear")
async def clear_ppct(
    grade: Optional[str] = None,
    subject: Optional[str] = None,
    user_id: str = Depends(get_current_user_id),
    db: Session = Depends(get_db)
):
    q = db.query(PPCTModel).filter(PPCTModel.user_id == user_id)
    if grade:
        q = q.filter(PPCTModel.grade == grade)
    if subject:
        q = q.filter(PPCTModel.subject == subject)
    count = q.delete()
    db.commit()
    return {"status": "success", "message": f"Đã xóa {count} bản ghi PPCT."}

from pydantic import BaseModel
class PPCTBulkDeleteRequest(BaseModel):
    ids: List[int]

@router.post("/bulk-delete")
async def bulk_delete_ppct(
    req: PPCTBulkDeleteRequest,
    user_id: str = Depends(get_current_user_id),
    db: Session = Depends(get_db)
):
    if not req.ids:
        return {"status": "success", "message": "Không có bản ghi nào được chọn"}
    count = db.query(PPCTModel).filter(
        PPCTModel.id.in_(req.ids),
        PPCTModel.user_id == user_id
    ).delete(synchronize_session=False)
    db.commit()
    return {"status": "success", "message": f"Đã xóa {count} tiết PPCT."}

@router.delete("/{item_id}")
async def delete_ppct_item(
    item_id: int,
    user_id: str = Depends(get_current_user_id),
    db: Session = Depends(get_db)
):
    item = db.query(PPCTModel).filter(PPCTModel.id == item_id, PPCTModel.user_id == user_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Bản ghi không tồn tại")
    db.delete(item)
    db.commit()
    return {"status": "success", "message": "Đã xóa tiết PPCT thành công."}

def format_ppct_lesson_str(lesson_number: Optional[int], notes: Optional[str] = "", title: Optional[str] = "") -> str:
    if not lesson_number:
        return ""
    is_cd = (notes and ("chuyên đề" in notes.lower() or "cđ" in notes.lower())) or \
            (title and ("chuyên đề" in title.lower() or "cđ" in title.lower())) or \
            (105 < lesson_number <= 150)
    if is_cd:
        cd_num = (lesson_number - 105) if lesson_number > 105 else lesson_number
        return f"{cd_num}CĐ"
    return str(lesson_number)

@router.get("/export/excel")
async def export_ppct_excel(
    grade: Optional[str] = Query(None),
    subject: Optional[str] = Query(None),
    user_id: str = Depends(get_current_user_id),
    db: Session = Depends(get_db)
):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from fastapi.responses import Response
    import io, urllib.parse

    q = db.query(PPCTModel).filter(PPCTModel.user_id == user_id)
    if grade and grade != "all":
        q = q.filter(PPCTModel.grade == grade)
    if subject and subject != "all":
        q = q.filter(PPCTModel.subject == subject)
    
    items = q.order_by(PPCTModel.week.asc(), PPCTModel.lesson_number.asc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PPCT"

    # Title
    ws.merge_cells("A1:F1")
    ws["A1"] = f"PHÂN PHỐI CHƯƠNG TRÌNH - {subject or 'TẤT CẢ MÔN'} ({grade or 'TẤT CẢ KHỐI'})"
    ws["A1"].font = Font(name="Times New Roman", size=14, bold=True, color="1F4E78")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = ["Tuần", "Tiết PPCT", "Khối", "Môn Học", "Tên Bài Dạy / Nội Dung Giảng Dạy", "Ghi Chú / ĐDDH"]
    ws.append(headers)

    header_font = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    for col_idx in range(1, 7):
        cell = ws.cell(row=2, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
    ws.row_dimensions[2].height = 25

    row_font = Font(name="Times New Roman", size=11)
    for idx, itm in enumerate(items, start=3):
        formatted_lesson = format_ppct_lesson_str(itm.lesson_number, itm.notes, itm.lesson_title)
        ws.cell(row=idx, column=1, value=itm.week).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=idx, column=2, value=formatted_lesson).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=idx, column=3, value=itm.grade).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=idx, column=4, value=itm.subject).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=idx, column=5, value=itm.lesson_title).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.cell(row=idx, column=6, value=itm.notes or "").alignment = Alignment(horizontal="left", vertical="center")

        for c in range(1, 7):
            cell = ws.cell(row=idx, column=c)
            cell.font = row_font
            cell.border = thin_border
            if idx % 2 == 0:
                cell.fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
        ws.row_dimensions[idx].height = 22

    # Column widths
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 45
    ws.column_dimensions["F"].width = 25

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"PPCT_{subject or 'Chung'}_{grade or 'Tat_Ca'}.xlsx"
    encoded_filename = urllib.parse.quote(filename)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"}
    )

class PPCTRawItem(BaseModel):
    grade: Optional[str] = "Khối 10"
    subject: Optional[str] = "Toán"
    week: int = 1
    lesson_number: int = 1
    lesson_title: str
    notes: Optional[str] = ""

class PPCTBulkCreateRequest(BaseModel):
    grade: str = "Khối 10"
    subject: str = "Toán"
    items: List[PPCTRawItem]
    overwrite: Optional[bool] = True

@router.post("/bulk", response_model=PPCTBulkUploadResponse)
async def bulk_create_ppct(
    req: PPCTBulkCreateRequest,
    user_id: str = Depends(get_current_user_id),
    db: Session = Depends(get_db)
):
    if req.overwrite:
        db.query(PPCTModel).filter(
            PPCTModel.user_id == user_id,
            PPCTModel.grade == req.grade,
            PPCTModel.subject == req.subject
        ).delete()
        db.commit()

    saved_items = []
    for itm in req.items:
        ppct_item = PPCTModel(
            user_id=user_id,
            grade=itm.grade or req.grade,
            subject=itm.subject or req.subject,
            week=itm.week,
            lesson_number=itm.lesson_number,
            lesson_title=itm.lesson_title,
            notes=itm.notes,
            semester="Học kỳ 1" if itm.week <= 18 else "Học kỳ 2"
        )
        db.add(ppct_item)
        saved_items.append(ppct_item)

    db.commit()
    for item in saved_items:
        db.refresh(item)

    return PPCTBulkUploadResponse(
        total_imported=len(saved_items),
        grade=req.grade,
        subject=req.subject,
        items=saved_items
    )

class PPCTPasteRequest(BaseModel):
    text: str
    grade: Optional[str] = "Khối 10"
    subject: Optional[str] = "Toán"
    overwrite: Optional[bool] = True
    api_key: Optional[str] = None
    model_name: Optional[str] = None


@router.post("/paste", response_model=PPCTBulkUploadResponse)
async def paste_ppct_text(
    req: PPCTPasteRequest,
    user_id: str = Depends(get_current_user_id),
    db: Session = Depends(get_db)
):
    from services.gemini_service import set_custom_gemini_key
    if req.api_key and req.api_key.strip():
        await run_in_threadpool(set_custom_gemini_key, req.api_key.strip())

    from services.ppct_parser import parse_ppct_text
    if not req.text or not req.text.strip():
        raise HTTPException(status_code=400, detail="Nội dung văn bản dán không được để trống")

    grade = normalize_grade(req.grade, req.text)
    subject = req.subject or "Toán"
    parsed_items = await run_in_threadpool(parse_ppct_text, req.text, default_grade=grade, default_subject=subject, selected_model=req.model_name)

    if not parsed_items:
        raise HTTPException(status_code=400, detail="Không nhận diện được tiết học từ nội dung đã dán.")

    if req.overwrite:
        db.query(PPCTModel).filter(
            PPCTModel.user_id == user_id,
            PPCTModel.grade == grade,
            PPCTModel.subject == subject
        ).delete()
        db.commit()

    saved_items = []
    for item_data in parsed_items:
        ppct_item = PPCTModel(
            user_id=user_id,
            grade=item_data.get("grade", grade),
            subject=item_data.get("subject", subject),
            week=item_data.get("week", 1),
            lesson_number=item_data.get("lesson_number", 1),
            lesson_title=item_data.get("lesson_title", "Bài học"),
            notes=item_data.get("notes"),
            semester=item_data.get("semester", "Học kỳ 1")
        )
        db.add(ppct_item)
        saved_items.append(ppct_item)

    db.commit()
    for item in saved_items:
        db.refresh(item)

    return PPCTBulkUploadResponse(
        total_imported=len(saved_items),
        grade=grade,
        subject=subject,
        items=saved_items
    )

@router.post("/upload-image", response_model=PPCTBulkUploadResponse)
async def upload_ppct_image(
    file: UploadFile = File(...),
    grade: str = Form("Khối 10"),
    subject: str = Form("Toán"),
    overwrite: bool = Form(True),
    api_key: Optional[str] = Form(None),
    model_name: Optional[str] = Form(None),
    user_id: str = Depends(get_current_user_id),
    db: Session = Depends(get_db)
):
    """
    Analyzes an uploaded or clipboard-pasted image of PPCT using Gemini AI Vision
    """
    from services.gemini_service import set_custom_gemini_key
    if api_key and api_key.strip():
        await run_in_threadpool(set_custom_gemini_key, api_key.strip())
    import json
    import re
    from services.gemini_service import generate_with_gemini_vision

    grade = normalize_grade(grade, file.filename or "")

    image_bytes = await file.read()
    if not image_bytes:
        raise HTTPException(status_code=400, detail="File ảnh rỗng hoặc không hợp lệ")

    prompt = f"""
Bạn là chuyên gia số hóa chương trình giáo dục phổ thông Việt Nam (Chuẩn GDPT 2018) bậc THPT.
Nhiệm vụ tối quan trọng: Đọc kỹ ảnh chụp bảng Phân phối chương trình (PPCT) môn {subject} {grade} và trích xuất TOÀN BỘ VÀ CHÍNH XÁC 100% danh sách từng tiết bài dạy.

================ CẤU TRÚC BẢNG ĐẶC TRƯNG CỦA PPCT TOÁN GDPT 2018 (BẢNG 5 CỘT) ================
Bảng gồm 5 cột chính:
- Cột 1: [Tuần] (ví dụ: '1 (7/9-12/9)', '2 (14/9-19/9)' -> Tuần lấy con số đầu: 1, 2, 3... 35).
- Cột 2: [Tiết] (Đánh số từ 1 đến 105 cho phần Chính khóa).
- Cột 3: [Chương/Bài] (Tên bài dạy Chính khóa).
- Cột 4: [Tiết] (Đánh số từ 1 đến 35 cho phần Chuyên đề học tập).
- Cột 5: [Chuyên đề học tập] (Tên bài Chuyên đề học tập).

================ QUY TẮC BẮT BUỘC ĐỂ TRÍCH XUẤT CHÍNH XÁC 100% ================

1. PHÂN TÁCH ĐỒNG THỜI CHÍNH KHÓA & CHUYÊN ĐỀ TRÊN TỪNG DÒNG:
   - Khi Cột 2 có số tiết: Trích xuất một tiết Chính khóa với:
     + "week": số tuần hiện tại (1..35)
     + "lesson_number": số nguyên ở Cột 2 (ví dụ: 1, 2, 3... 105)
     + "lesson_title": nội dung chữ ở Cột 3 (ví dụ: "§1. Mệnh đề", "§2. Tập hợp và các phép toán trên tập hợp")
     + "notes": ""
   - Khi Cột 4 có số tiết: Trích xuất một tiết Chuyên đề với:
     + "week": số tuần hiện tại (1..35)
     + "lesson_number": 105 + số nguyên ở Cột 4 (ví dụ: Tiết CĐ 1 ghi 106, Tiết CĐ 2 ghi 107, Tiết CĐ 3 ghi 108... Tiết CĐ 35 ghi 140)
     + "lesson_title": nội dung chữ ở Cột 5 (ví dụ: "CĐ1-§1. Hệ phương trình bậc nhất ba ẩn", "CĐ1-§2. Ứng dụng của hệ phương trình bậc nhất ba ẩn")
     + "notes": "Chuyên đề"

2. QUY TẮC Ô GỘP VÀ Ô TRỐNG (MERGED CELLS):
   - Cột 1 (Tuần) thường bị gộp cho 3-4 dòng: Mọi dòng nằm trong khối gộp đó đều thuộc về Tuần đó cho đến khi xuất hiện số Tuần tiếp theo.
   - Cột 4 & Cột 5 (Chuyên đề) thường chỉ có 1 tiết/tuần ở dòng đầu của tuần đó, các dòng sau bị trống -> Chỉ tạo tiết Chuyên đề khi Cột 4 có số tiết, TUYỆT ĐỐI KHÔNG tạo tiết Chuyên đề ở các dòng trống.

3. LOẠI BỎ CÁC DÒNG TIÊU ĐỀ CHƯƠNG / CHUYÊN ĐỀ TỔNG QUÁT:
   - Các dòng tiêu đề in hoa không có số tiết ở Cột 2 và Cột 4 (như:
     "CHƯƠNG I. MỆNH ĐỀ VÀ TẬP HỢP (9 TIẾT)",
     "CHUYÊN ĐỀ 1. HỆ PHƯƠNG TRÌNH BẬC NHẤT BA ẨN (11 TIẾT)",
     "CHƯƠNG II. BẤT PHƯƠNG TRÌNH VÀ HỆ BẤT PHƯƠNG TRÌNH BẬC NHẤT HAI ẨN (6 TIẾT)",
     "CHƯƠNG III. HỆ THỨC LƯỢNG TRONG TAM GIÁC (7 TIẾT)")
     -> ĐÂY LÀ DÒNG BANNER TIÊU ĐỀ CHƯƠNG, BỎ QUA HOÀN TOÀN, KHÔNG ĐƯỢC TẠO TIẾT HỌC CHO CÁC DÒNG NÀY!

4. GIỮ NGUYÊN VĂN TÊN BÀI HỌC THEO BẢNG GỐC:
   - Nếu bảng gốc ghi từng dòng lặp lại:
     Dòng 1: Tiết 1 | §1. Mệnh đề
     Dòng 2: Tiết 2 | §1. Mệnh đề
     Dòng 3: Tiết 3 | §1. Mệnh đề
     Dòng 4: Tiết 4 | §1. Mệnh đề
     -> Giữ nguyên tên bài "§1. Mệnh đề" cho cả 4 tiết với lesson_number tương ứng 1, 2, 3, 4. Không tự ý thêm bớt hậu tố nếu bảng đã chia từng dòng riêng.

5. VÍ DỤ MINH HỌA ĐẦU RA MẪU:
[
  {{ "week": 1, "lesson_number": 1, "lesson_title": "§1. Mệnh đề", "notes": "" }},
  {{ "week": 1, "lesson_number": 2, "lesson_title": "§1. Mệnh đề", "notes": "" }},
  {{ "week": 1, "lesson_number": 3, "lesson_title": "§1. Mệnh đề", "notes": "" }},
  {{ "week": 1, "lesson_number": 106, "lesson_title": "CĐ1-§1. Hệ phương trình bậc nhất ba ẩn", "notes": "Chuyên đề" }},
  {{ "week": 2, "lesson_number": 4, "lesson_title": "§1. Mệnh đề", "notes": "" }},
  {{ "week": 2, "lesson_number": 5, "lesson_title": "§2. Tập hợp và các phép toán trên tập hợp", "notes": "" }},
  {{ "week": 2, "lesson_number": 6, "lesson_title": "§2. Tập hợp và các phép toán trên tập hợp", "notes": "" }},
  {{ "week": 2, "lesson_number": 107, "lesson_title": "CĐ1-§1. Hệ phương trình bậc nhất ba ẩn", "notes": "Chuyên đề" }},
  {{ "week": 3, "lesson_number": 7, "lesson_title": "§2. Tập hợp và các phép toán trên tập hợp", "notes": "" }},
  {{ "week": 3, "lesson_number": 8, "lesson_title": "§2. Tập hợp và các phép toán trên tập hợp", "notes": "" }},
  {{ "week": 3, "lesson_number": 9, "lesson_title": "Bài tập cuối chương I", "notes": "" }},
  {{ "week": 3, "lesson_number": 108, "lesson_title": "CĐ1-§1. Hệ phương trình bậc nhất ba ẩn", "notes": "Chuyên đề" }}
]

TRẢ VỀ DUY NHẤT 1 MẢNG JSON HỢP LỆ (bắt đầu bằng [ và kết thúc bằng ]):
"""
    try:
        raw_result = await run_in_threadpool(
            generate_with_gemini_vision,
            prompt=prompt,
            image_bytes=image_bytes,
            mime_type=file.content_type or "image/png",
            user_key=api_key,
            selected_model=model_name
        )
    except Exception as e:
        logger.error(f"Gemini Vision PPCT extraction error: {e}")
        raise HTTPException(status_code=500, detail=f"Lỗi AI Vision khi đọc ảnh: {str(e)}")

    # Clean JSON
    cleaned = raw_result.strip()
    if cleaned.startswith("```json"):
        cleaned = cleaned[7:]
    if cleaned.startswith("```"):
        cleaned = cleaned[3:]
    if cleaned.endswith("```"):
        cleaned = cleaned[:-3]
    cleaned = cleaned.strip()

    try:
        parsed_items = json.loads(cleaned)
    except Exception as e:
        logger.warning(f"Could not parse direct JSON from Vision: {cleaned}")
        match = re.search(r'\[.*\]', cleaned, re.DOTALL)
        if match:
            parsed_items = json.loads(match.group(0))
        else:
            raise HTTPException(status_code=500, detail="AI không thể tạo cấu trúc bảng từ ảnh này. Vui lòng thử ảnh chụp rõ nét hơn.")

    if not isinstance(parsed_items, list) or len(parsed_items) == 0:
        raise HTTPException(status_code=400, detail="Không tìm thấy dữ liệu tiết học trong ảnh.")

    from services.ppct_parser import is_valid_lesson_title

    if overwrite:
        db.query(PPCTModel).filter(
            PPCTModel.user_id == user_id,
            PPCTModel.grade == grade,
            PPCTModel.subject == subject
        ).delete()
        db.commit()

    saved_items = []
    for item_data in parsed_items:
        title = str(item_data.get("lesson_title") or "Bài học").strip()
        if not is_valid_lesson_title(title):
            continue

        w = int(item_data.get("week") or 1)
        l_num = int(item_data.get("lesson_number") or len(saved_items) + 1)
        notes = str(item_data.get("notes") or "").strip()

        is_cd = ("chuyên đề" in title.lower() or "cđ" in title.lower() or "chuyên đề" in notes.lower() or l_num > 105)
        if is_cd:
            if "chuyên đề" not in notes.lower():
                notes = f"Chuyên đề{(' - ' + notes) if notes else ''}"
            if l_num <= 35:
                l_num = 105 + l_num

        ppct_item = PPCTModel(
            user_id=user_id,
            grade=grade,
            subject=subject,
            week=w,
            lesson_number=l_num,
            lesson_title=title,
            notes=notes,
            semester="Học kỳ 1" if w <= 18 else "Học kỳ 2"
        )
        db.add(ppct_item)
        saved_items.append(ppct_item)

    db.commit()
    for item in saved_items:
        db.refresh(item)

    return PPCTBulkUploadResponse(
        total_imported=len(saved_items),
        grade=grade,
        subject=subject,
        items=saved_items
    )


